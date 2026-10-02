import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import subprocess
import os
import pypdf

doc = docx.Document()

# Set A4 margins
section = doc.sections[0]
section.page_width = Inches(8.27)   # 210 mm
section.page_height = Inches(11.69) # 297 mm
section.top_margin = Inches(0.52)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)

style = doc.styles['Normal']
font = style.font
font.name = 'Liberation Serif'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0, 0, 0)

def add_header(name, email, phone, linkedin_text):
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(name)
    r_name.font.name = 'Liberation Serif'
    r_name.font.size = Pt(22)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(6)
    p_contact.paragraph_format.line_spacing = 1.0
    
    parts = []
    if email: parts.append(email)
    if phone: parts.append(phone)
    contact_text = " | ".join(parts)
    if contact_text:
        r = p_contact.add_run(contact_text + " | ")
        r.font.name = 'Liberation Serif'
        r.font.size = Pt(10)
    
    r_li = p_contact.add_run(linkedin_text)
    r_li.font.name = 'Liberation Serif'
    r_li.font.size = Pt(10)
    r_li.font.underline = True

def add_section_title(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Liberation Serif'
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.underline = True

def add_entry_heading(org, location, role, date):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(5.25), Inches(1.86)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcMar = OxmlElement('w:tcMar')
            for m in ['top', 'bottom', 'left', 'right']:
                node = OxmlElement(f'w:{m}')
                node.set(qn('w:w'), '0')
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)
            
    # Row 0: Org (bold) + Location (bold, right)
    p00 = table.rows[0].cells[0].paragraphs[0]
    p00.paragraph_format.space_before = Pt(2)
    p00.paragraph_format.space_after = Pt(0)
    p00.paragraph_format.line_spacing = 1.03
    r00 = p00.add_run(org)
    r00.font.name = 'Liberation Serif'
    r00.font.size = Pt(10.5)
    r00.font.bold = True

    p01 = table.rows[0].cells[1].paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p01.paragraph_format.space_before = Pt(2)
    p01.paragraph_format.space_after = Pt(0)
    p01.paragraph_format.line_spacing = 1.03
    r01 = p01.add_run(location)
    r01.font.name = 'Liberation Serif'
    r01.font.size = Pt(10.5)
    r01.font.bold = True

    # Row 1: Role (italic) + Date (regular, right)
    p10 = table.rows[1].cells[0].paragraphs[0]
    p10.paragraph_format.space_before = Pt(0)
    p10.paragraph_format.space_after = Pt(0.5)
    p10.paragraph_format.line_spacing = 1.03
    r10 = p10.add_run(role)
    r10.font.name = 'Liberation Serif'
    r10.font.size = Pt(10)
    r10.font.italic = True

    p11 = table.rows[1].cells[1].paragraphs[0]
    p11.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p11.paragraph_format.space_before = Pt(0)
    p11.paragraph_format.space_after = Pt(0.5)
    p11.paragraph_format.line_spacing = 1.03
    r11 = p11.add_run(date)
    r11.font.name = 'Liberation Serif'
    r11.font.size = Pt(10)
    r11.font.italic = False

def add_bullet(text, bold_prefix=None, space_after=Pt(1.0)):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.04
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = 'Liberation Serif'
    r_bullet.font.size = Pt(10)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Liberation Serif'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = 'Liberation Serif'
    r_text.font.size = Pt(10)

def add_subheading(text, space_before=Pt(2.5), space_after=Pt(0.5)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run(text)
    r.font.name = 'Liberation Serif'
    r.font.size = Pt(10)
    r.font.bold = True

def add_skill_line(label, content, space_after=Pt(1.5)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.04
    
    r_lbl = p.add_run(label + ": ")
    r_lbl.font.name = 'Liberation Serif'
    r_lbl.font.size = Pt(10)
    r_lbl.font.italic = True
    
    r_cnt = p.add_run(content)
    r_cnt.font.name = 'Liberation Serif'
    r_cnt.font.size = Pt(10)

# Header
add_header("Alexandru Mihalcea-Calinescu", "alex.mihalcea12@gmail.com", "+40 722 237 081", "LinkedIn")

# EDUCATION
add_section_title("EDUCATION")
add_entry_heading("Università Bocconi", "Milan, Italy", "BSc in Economics, Management and Computer Science", "Sep 2026 - Present")
add_bullet("Student Societies: Member of BSML (Machine Learning) and Hephaestus (Applied AI Association).")

add_entry_heading("Avenor College", "Bucharest, Romania", "A Levels, British Curriculum", "Jun 2026")
add_bullet("Economics (A), Mathematics (A), Computer Science (A), History (A), Extended Project Qualification (A)")
add_bullet("IGCSE results: A*A*A*A*A*AAA(A)")
add_bullet("SAT: 1500, 99th percentile (800 Reading and Writing, 700 Mathematics)")

# AWARDS & ACHIEVEMENTS
add_section_title("AWARDS & ACHIEVEMENTS")
add_bullet("continental top-three team and individual Silver Award at the International Final among 428 teams and 2,502 competitors from 47 countries and regions", bold_prefix="World Economics Cup 2025: ")
add_bullet("Silver Award", bold_prefix="UKMT Senior Mathematical Challenge: ")
add_bullet("Silver and Bronze", bold_prefix="Duke of Edinburgh's International Award: ")
add_bullet("4+ years", bold_prefix="Academic Honour Roll: ")

# WORK EXPERIENCE & PROJECTS
add_section_title("WORK EXPERIENCE & PROJECTS")
add_entry_heading("PwC Romania", "Bucharest, Romania", "M&A Intern - Corporate Finance", "Oct 2025")
add_bullet("Researched 20+ companies across multiple industries, examining business models, market positioning and financial performance")
add_bullet("Synthesised company research and analysis of trading comparables, precedent transaction multiples, historical trading data and financial statements into PowerPoint presentations for senior Corporate Finance team members")

add_entry_heading("BAINSA Hackathon 2026", "Milan, Italy", "FocusAid - Applied AI & Full-Stack Contributor", "Sep 2026")
add_bullet("Engineered an AI accessibility web app for neurodivergent students with real-time lecture catch-up and AI summaries")
add_bullet("Integrated private client-side attention tracking via MediaPipe and on-device WebGPU LLM inference")

add_entry_heading("Genpact & IB Cargo", "Bucharest, Romania", "Operations & Logistics Insight Programmes", "Jun 2024 - Jun 2025")
add_bullet("Reviewed operational workflows across order-to-cash, accounts payable, payroll, contract administration, and international supply-chain logistics through cross-departmental sessions with senior leaders")

add_entry_heading("Eau de Web", "Bucharest, Romania", "Summer Intern - Software Development", "Jun 2023")
add_bullet("Contributed to front-end development of a university-advisory platform, implementing UI components using HTML and CSS")

# EXTRACURRICULAR ACTIVITIES
add_section_title("EXTRACURRICULAR ACTIVITIES")
add_entry_heading("Christmas Fair Organisation Team", "Bucharest, Romania", "Volunteer, Communications & Finance Coordinator", "2022 - 2024")
add_bullet("Volunteered at one fair; coordinated communications and sponsorship outreach for 4+ months the following year; managed budgeting and financial administration during a subsequent 4+ month cycle")

add_subheading("Additional Activities")
add_bullet("Executive production for TEDx Avenor (2023); sponsorship team member for a 24-hour business challenge")
add_bullet("Community initiatives: Baneasa Forest campaign (2023-2024), Avenor Forest Run (2022), Time to Play NGO mobility project (2025)")

# AUTHORED WORKS
add_section_title("AUTHORED WORKS")
add_bullet("An Investigation into Existing and Potential Economic Impacts of Internet Restrictions (EPQ)")
add_bullet("To What Extent Can Playing an Educational Game Help People Struggling with Dyslexia? (.XIA Software Project)")

# SKILLS & INTERESTS
add_section_title("SKILLS & INTERESTS")
add_skill_line("Languages", "Romanian (Native), English (Fluent; IELTS 8.5), Spanish (DELE B1), Portuguese (Intermediate), Italian (Basic)")
add_skill_line("Technical Skills", "Python, TypeScript, Next.js, React, WebGPU, MediaPipe, Git, HTML/CSS, Excel, PowerPoint")
add_skill_line("Interests", "Table football, table tennis, custom PC building, open-source OS, European history")

output_docx = "/Users/alexmih/portfolio/docs/Alexandru_Mihalcea_Calinescu_CV.docx"
output_pdf = "/Users/alexmih/portfolio/docs/Alexandru_Mihalcea_Calinescu_CV.pdf"
doc.save(output_docx)

subprocess.run([
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "--headless",
    "--convert-to", "pdf",
    output_docx,
    "--outdir", "/Users/alexmih/portfolio/docs"
], check=True)

reader = pypdf.PdfReader(output_pdf)
print(f"Generated PDF Page Count: {len(reader.pages)}")
